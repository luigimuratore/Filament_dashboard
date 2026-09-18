"""Workbook operations shared by the dashboard and its tests."""
from collections import defaultdict
from datetime import datetime, time, timedelta
from pathlib import Path
from uuid import uuid4
from filament_lock import file_lock
import hashlib
import math
import os
import shutil
import tempfile
from openpyxl import load_workbook

FILE = Path(__file__).resolve().parent / 'Tracker_Filament_Dashboard.xlsx'
PLANNING_SHEET = 'Pianificazione'
PLANNING_PRIORITIES = ('SUBITO', 'Urgente', 'Quando possibile')
PLANNING_PRIORITY_RANK = {'SUBITO': 3, 'Urgente': 2, 'Quando possibile': 1}
PRINT_METADATA_HEADERS = {
    8: 'ID Stampa', 9: 'Durata (min)', 10: 'Inizio nel calendario',
}
PLANNING_HEADERS = [
    'ID Pianificazione', 'Nome stampa', 'Durata (min)', 'Inizio previsto',
    'Ugello', 'ID Bobina', 'Materiale', 'Grammi previsti', 'Note', 'Creata il',
    'Priorità',
]


def load(path=FILE):
    # Read the bytes once so the revision always matches the loaded workbook.
    from io import BytesIO
    data = Path(path).read_bytes()
    wb = load_workbook(BytesIO(data))
    wb._revision = hashlib.sha256(data).digest()
    return wb


def save(wb, path=FILE):
    path = Path(path)
    with file_lock(path.with_suffix('.lock')):
        if hashlib.sha256(path.read_bytes()).digest() != wb._revision:
            raise ValueError('I dati sono cambiati in un’altra finestra. Ricarica la pagina e riprova.')
        fd, name = tempfile.mkstemp(dir=path.parent, suffix='.xlsx')
        os.close(fd)
        try:
            wb.save(name)
            shutil.copy2(path, path.with_suffix('.backup.xlsx'))
            os.replace(name, path)
        finally:
            if os.path.exists(name):
                os.unlink(name)


def assignments(wb):
    return {int(r[0]): str(r[1] or '') for r in wb['Ugelli'].iter_rows(min_row=2, values_only=True) if r[0] in (1, 2, 3)}


def bobine(wb, include_removed=False):
    loaded = {bid: n for n, bid in assignments(wb).items() if bid}
    out = []
    for row in wb['Bobine'].iter_rows(min_row=2):
        if not row[0].value:
            continue
        vals = [c.value for c in row]
        removed = len(vals) > 10 and vals[10] == 'Eliminata'
        if removed and not include_removed:
            continue
        weight, used = float(vals[4] or 0), float(vals[5] or 0)
        remaining = max(0, weight - used)
        bid = str(vals[0])
        out.append(dict(row=row[0].row, id=bid, eliminata=removed, materiale=str(vals[1] or ''), marca=str(vals[2] or ''),
                        colore=str(vals[3] or ''), peso=weight, usati=used, rim=remaining,
                        pct=remaining / weight * 100 if weight else 0, ugello=loaded.get(bid),
                        stato=f'Ugello {loaded[bid]}' if bid in loaded else ('Esaurita' if remaining == 0 else 'Magazzino')))
    return out


def health(pct):
    if pct <= 0: return 'Esaurita', 'red'
    if pct <= 20: return 'Scorta critica', 'red'
    if pct <= 40: return 'In diminuzione', 'amber'
    return 'Buona disponibilità', 'green'


def sync_states(wb):
    for b in bobine(wb):
        wb['Bobine'].cell(b['row'], 9, 'Esaurita' if not b['rim'] else ('Caricata' if b['ugello'] else 'Disponibile'))


def set_assignments(wb, choices):
    ids = [x for x in choices.values() if x]
    if len(ids) != len(set(ids)):
        raise ValueError('Ogni bobina può essere caricata su un solo ugello. Controlla le selezioni.')
    bs = {b['id']: b for b in bobine(wb)}
    for bid in ids:
        if bid not in bs or bs[bid]['rim'] <= 0:
            raise ValueError(f'{bid}: scegli una bobina con materiale disponibile.')
    for n in (1, 2, 3):
        wb['Ugelli'].cell(n + 1, 2).value = choices.get(n, '')
    sync_states(wb)


def _duration(value, required=False):
    if value is None and not required:
        return None
    try:
        minutes = int(value)
    except (ValueError, TypeError, OverflowError):
        raise ValueError('Inserisci una durata valida in ore e minuti.')
    if minutes <= 0:
        raise ValueError('La durata deve essere maggiore di zero.')
    return minutes


def planning_priority(value):
    if value in (None, ''):
        return 'Quando possibile'
    normalized = str(value).strip().casefold()
    choices = {priority.casefold(): priority for priority in PLANNING_PRIORITIES}
    if normalized not in choices:
        raise ValueError('Scegli una priorità tra SUBITO, Urgente e Quando possibile.')
    return choices[normalized]


def record_print(wb, name, consumption, note='', when=None, duration_minutes=None):
    if not name.strip(): raise ValueError('Inserisci il nome della stampa.')
    if any(not math.isfinite(g) or g < 0 for g in consumption.values()):
        raise ValueError('I consumi devono essere numeri positivi o zero.')
    if sum(consumption.values()) <= 0: raise ValueError('Inserisci un consumo maggiore di zero per almeno un ugello.')
    ass = assignments(wb)
    bids = [x for x in ass.values() if x]
    if len(bids) != len(set(bids)):
        raise ValueError('La stessa bobina è assegnata a più ugelli. Correggi il setup prima di registrare.')
    bs = {b['id']: b for b in bobine(wb)}
    for n, g in consumption.items():
        if not g: continue
        b = bs.get(ass.get(n))
        if not b: raise ValueError(f'Ugello {n}: carica una bobina prima di inserire il consumo.')
        if g > b['rim']: raise ValueError(f"Ugello {n} · {b['id']}: richiesti {g:g} g, disponibili {b['rim']:g} g.")
    duration_minutes = _duration(duration_minutes)
    now, pid = when or datetime.now(), str(uuid4())
    ws = wb['Stampe']
    for column, heading in PRINT_METADATA_HEADERS.items():
        ws.cell(1, column, heading)
    for n, g in consumption.items():
        if g <= 0: continue
        b = bs[ass[n]]
        wb['Bobine'].cell(b['row'], 6, b['usati'] + g)
        ws.append([now, name.strip(), n, b['id'], b['materiale'], g, note.strip(), pid, duration_minutes, None])
    sync_states(wb)


def history(wb):
    groups = {}
    for row_number, row in enumerate(wb['Stampe'].iter_rows(min_row=2, values_only=True), 2):
        if not row[0] and not row[1]: continue
        raw = row[0]
        if isinstance(raw, datetime): date = raw
        else:
            try: date = datetime.fromisoformat(str(raw))
            except (ValueError, TypeError): date = datetime.min
        key = row[7] if len(row) > 7 and row[7] else (str(raw), row[1], row[6])
        duration = None
        if len(row) > 8 and row[8] not in (None, ''):
            try: duration = int(row[8])
            except (ValueError, TypeError): pass
        raw_planned_start = row[9] if len(row) > 9 else None
        if isinstance(raw_planned_start, datetime):
            planned_start = raw_planned_start
        elif raw_planned_start:
            try: planned_start = datetime.fromisoformat(str(raw_planned_start))
            except (ValueError, TypeError): planned_start = None
        else:
            planned_start = None
        item = groups.setdefault(key, dict(
            data=date, nome=str(row[1] or ''), note=str(row[6] or ''),
            durata=duration, consumi=[], totale=0.0, rows=[],
            key=str(row[7]) if len(row) > 7 and row[7] else f'legacy-{row_number}',
            planned_start=planned_start, calendar_start=planned_start or date,
        ))
        if item['planned_start'] is None and planned_start is not None:
            item['planned_start'] = planned_start
            item['calendar_start'] = planned_start
        item['rows'].append(row_number)
        grams = float(row[5] or 0)
        item['consumi'].append(dict(ugello=row[2], bobina=row[3], materiale=row[4], grammi=grams))
        item['totale'] += grams
    return sorted(groups.values(), key=lambda x: x['data'], reverse=True)


def _planning_sheet(wb):
    if PLANNING_SHEET not in wb.sheetnames:
        ws = wb.create_sheet(PLANNING_SHEET)
        ws.append(PLANNING_HEADERS)
        ws.freeze_panes = 'A2'
        return ws
    ws = wb[PLANNING_SHEET]
    for column, heading in enumerate(PLANNING_HEADERS, 1):
        ws.cell(1, column, heading)
    return ws


def planned_prints(wb):
    if PLANNING_SHEET not in wb.sheetnames:
        return []
    groups = {}
    for row_number, row in enumerate(wb[PLANNING_SHEET].iter_rows(min_row=2, values_only=True), 2):
        if not row[0]:
            continue
        key = str(row[0])
        raw_start = row[3]
        if isinstance(raw_start, datetime):
            start = raw_start
        elif raw_start:
            try: start = datetime.fromisoformat(str(raw_start))
            except (ValueError, TypeError): start = None
        else:
            start = None
        raw_created = row[9] if len(row) > 9 else None
        if isinstance(raw_created, datetime):
            created = raw_created
        else:
            try: created = datetime.fromisoformat(str(raw_created))
            except (ValueError, TypeError): created = datetime.min
        try: duration = int(row[2])
        except (ValueError, TypeError): duration = 0
        priority = planning_priority(row[10] if len(row) > 10 else None)
        item = groups.setdefault(key, dict(
            key=key, nome=str(row[1] or ''), durata=duration, inizio=start,
            note=str(row[8] or ''), creata=created, consumi=[], totale=0.0,
            rows=[], priorita=priority,
        ))
        item['rows'].append(row_number)
        grams = float(row[7] or 0)
        item['consumi'].append(dict(
            ugello=int(row[4]), bobina=str(row[5] or ''),
            materiale=str(row[6] or ''), grammi=grams,
        ))
        item['totale'] += grams
    return sorted(groups.values(), key=lambda item: (
        (0, item['inizio'], item['creata'], item['nome'].casefold())
        if item['inizio'] is not None else
        (1, -PLANNING_PRIORITY_RANK[item['priorita']], item['creata'], item['nome'].casefold())
    ))


def suggest_next_print(items, now=None, work_start=time(8, 30), work_end=time(17, 30),
                       turnaround_minutes=30, workdays=(0, 1, 2, 3, 4)):
    now = now or datetime.now()
    scheduled = [item for item in items if item['inizio'] is not None]
    active = next((item for item in scheduled
                   if item['inizio'] <= now < item['inizio'] + timedelta(minutes=item['durata'])), None)
    if active is None:
        return None
    active_end = active['inizio'] + timedelta(minutes=active['durata'])
    future = sorted(
        (item for item in scheduled
         if item['key'] != active['key'] and item['inizio'] >= active_end),
        key=lambda item: item['inizio'],
    )
    queued = [item for item in items if item['inizio'] is None and item['durata'] > 0]
    timeline = [active] + future
    for index, previous in enumerate(timeline):
        start = previous['inizio'] + timedelta(minutes=previous['durata'] + turnaround_minutes)
        next_scheduled = timeline[index + 1] if index + 1 < len(timeline) else None
        latest_end = (next_scheduled['inizio'] - timedelta(minutes=turnaround_minutes)
                      if next_scheduled else None)
        candidates = []
        for item in queued:
            end = start + timedelta(minutes=item['durata'])
            if end.weekday() not in workdays or not work_start <= end.time() <= work_end:
                continue
            if latest_end is not None and end > latest_end:
                continue
            candidates.append(item)
        if not candidates:
            continue
        candidate = min(candidates, key=lambda item: (
            -PLANNING_PRIORITY_RANK[planning_priority(item.get('priorita'))],
            -item['durata'], item['creata'], item['nome'].casefold(), item['key'],
        ))
        end = start + timedelta(minutes=candidate['durata'])
        return {
            'attiva': active,
            'precedente': previous,
            'proposta': candidate,
            'inizio': start,
            'fine': end,
            'prossima_programmata': next_scheduled,
            'cambio_minuti': turnaround_minutes,
        }
    return None


def add_planned_print(wb, name, consumption, duration_minutes, note='', start=None, created=None,
                      priority='Quando possibile'):
    if not name.strip():
        raise ValueError('Inserisci il nome della stampa.')
    duration_minutes = _duration(duration_minutes, required=True)
    if start is not None and not isinstance(start, datetime):
        raise ValueError('Inserisci una data e un’ora valide.')
    if any(not math.isfinite(g) or g < 0 for g in consumption.values()):
        raise ValueError('I consumi previsti devono essere numeri positivi o zero.')
    if sum(consumption.values()) <= 0:
        raise ValueError('Inserisci un consumo previsto maggiore di zero per almeno un ugello.')
    ass = assignments(wb)
    bids = [x for x in ass.values() if x]
    if len(bids) != len(set(bids)):
        raise ValueError('La stessa bobina è assegnata a più ugelli. Correggi il setup prima di programmare.')
    bs = {b['id']: b for b in bobine(wb)}
    uses = []
    for n, grams in consumption.items():
        if grams <= 0:
            continue
        b = bs.get(ass.get(n))
        if not b:
            raise ValueError(f'Ugello {n}: carica una bobina prima di inserire il consumo previsto.')
        if grams > b['rim']:
            raise ValueError(f"Ugello {n} · {b['id']}: richiesti {grams:g} g, disponibili {b['rim']:g} g.")
        uses.append((n, b, grams))
    priority = planning_priority(priority)
    pid, created = str(uuid4()), created or datetime.now()
    if start is not None:
        _check_schedule_overlap(wb, start, duration_minutes)
    ws = _planning_sheet(wb)
    for n, b, grams in uses:
        ws.append([
            pid, name.strip(), duration_minutes, start, n, b['id'],
            b['materiale'], grams, note.strip(), created, priority,
        ])
    return pid


def _planned(wb, key):
    item = next((p for p in planned_prints(wb) if p['key'] == key), None)
    if not item:
        raise ValueError('Stampa pianificata non trovata. Ricarica la pagina.')
    return item


def _check_schedule_overlap(wb, start, duration_minutes, exclude=None):
    end = start + timedelta(minutes=duration_minutes)
    for other in planned_prints(wb):
        if other['key'] == exclude or other['inizio'] is None:
            continue
        other_end = other['inizio'] + timedelta(minutes=other['durata'])
        if start < other_end and end > other['inizio']:
            label = other['inizio'].strftime('%d/%m %H:%M')
            raise ValueError(f'L’orario si sovrappone a «{other["nome"]}» ({label}). Scegli uno spazio libero.')


def schedule_planned_print(wb, key, start, duration_minutes=None):
    item = _planned(wb, key)
    if not isinstance(start, datetime):
        raise ValueError('Inserisci una data e un’ora valide.')
    duration_minutes = _duration(duration_minutes if duration_minutes is not None else item['durata'], required=True)
    _check_schedule_overlap(wb, start, duration_minutes, exclude=key)
    ws = wb[PLANNING_SHEET]
    for row in item['rows']:
        ws.cell(row, 3, duration_minutes)
        ws.cell(row, 4, start)


def unschedule_planned_print(wb, key):
    item = _planned(wb, key)
    ws = wb[PLANNING_SHEET]
    for row in item['rows']:
        ws.cell(row, 4).value = None


def update_planned_priority(wb, key, priority):
    item = _planned(wb, key)
    priority = planning_priority(priority)
    ws = _planning_sheet(wb)
    for row in item['rows']:
        ws.cell(row, 11, priority)


def delete_planned_print(wb, key):
    item = _planned(wb, key)
    ws = wb[PLANNING_SHEET]
    for row in item['rows']:
        for cell in ws[row]:
            cell.value = None


def complete_planned_print(wb, key, when=None, consumption=None):
    item = _planned(wb, key)
    when = when or datetime.now()
    if not isinstance(when, datetime):
        raise ValueError('Inserisci una data e un’ora valide.')
    rows = item['consumi'] if consumption is None else consumption
    bs = {b['id']: b for b in bobine(wb, include_removed=True)}
    uses, requested = [], defaultdict(float)
    for use in rows:
        try:
            nozzle = int(use['ugello'])
            bid = str(use['bobina'])
            grams = float(use['grammi'])
        except (KeyError, TypeError, ValueError, OverflowError):
            raise ValueError('Controlla ugello, bobina e grammi di ogni riga.')
        if nozzle not in (1, 2, 3):
            raise ValueError('Scegli un ugello tra 1, 2 e 3.')
        if not math.isfinite(grams) or grams < 0:
            raise ValueError('I grammi devono essere positivi o zero.')
        if grams == 0:
            continue
        spool = bs.get(bid)
        if not spool or spool['eliminata']:
            raise ValueError(f'Bobina {bid} non disponibile: correggi il piano prima di completare la stampa.')
        requested[bid] += grams
        uses.append((nozzle, spool, grams))
    if not uses:
        raise ValueError('La stampa deve contenere almeno un consumo maggiore di zero.')
    for bid, grams in requested.items():
        if grams > bs[bid]['rim']:
            raise ValueError(f'{bid}: richiesti {grams:g} g, disponibili {bs[bid]["rim"]:g} g.')
    for bid, grams in requested.items():
        spool = bs[bid]
        wb['Bobine'].cell(spool['row'], 6, spool['usati'] + grams)
    pid = str(uuid4())
    ws = wb['Stampe']
    for column, heading in PRINT_METADATA_HEADERS.items():
        ws.cell(1, column, heading)
    # Keep the original slot after completion, even when the actual timestamp
    # is corrected in the confirmation dialog or later in the history editor.
    calendar_start = item['inizio'] or when
    for nozzle, spool, grams in uses:
        ws.append([
            when, item['nome'], nozzle, spool['id'], spool['materiale'], grams,
            item['note'], pid, item['durata'], calendar_start,
        ])
    delete_planned_print(wb, key)
    sync_states(wb)


def add_spool(wb, material, brand, color, weight):
    if not material.strip() or not brand.strip() or not color.strip():
        raise ValueError('Inserisci materiale, marca e colore della bobina.')
    if not math.isfinite(weight) or weight <= 0: raise ValueError('Il peso netto deve essere maggiore di zero.')
    bs = bobine(wb, include_removed=True)
    ids = [int(b['id'][1:]) for b in bs if b['id'].startswith('B') and b['id'][1:].isdigit()]
    bid = f'B{max(ids, default=0) + 1:03d}'
    r = max([b['row'] for b in bs], default=1) + 1
    values = [bid, material.strip(), brand.strip(), color.strip(), weight, 0,
              f'=MAX(0,E{r}-F{r})', f'=IFERROR(G{r}/E{r},0)', 'Disponibile']
    for c, v in enumerate(values, 1): wb['Bobine'].cell(r, c, v)
    wb['Bobine'].cell(r, 8).number_format = '0%'
    return bid


def restock(bs):
    groups = defaultdict(list)
    for b in bs: groups[(b['materiale'].casefold(), b['colore'].casefold())].append(b)
    out = []
    for group in groups.values():
        # Use one reference spool, so empty historic spools do not inflate the threshold.
        reference = max(b['peso'] for b in group)
        remaining = sum(b['rim'] for b in group)
        if remaining <= reference * .2:
            out.append({'Materiale': group[0]['materiale'], 'Colore': group[0]['colore'],
                        'Disponibili (g)': remaining, 'Soglia (g)': reference * .2,
                        'Ricambi in magazzino': sum(b['rim'] > 0 and not b['ugello'] for b in group),
                        'Priorità': 'Esaurito' if remaining == 0 else 'Da acquistare'})
    return sorted(out, key=lambda x: x['Disponibili (g)'])


def remove_spool(wb, bid):
    spool = next((b for b in bobine(wb) if b['id'] == bid), None)
    if spool is None:
        raise ValueError('Bobina non trovata nel magazzino.')
    if bid in assignments(wb).values():
        raise ValueError('Scarica prima la bobina dall’ugello nella Panoramica.')
    # Keep the row and ID so print history retains color/brand and IDs are never reused.
    wb['Bobine'].cell(1, 11, 'Archivio')
    wb['Bobine'].cell(spool['row'], 11, 'Eliminata')


def restore_spool(wb, bid):
    spool = next((b for b in bobine(wb, include_removed=True) if b['id'] == bid and b['eliminata']), None)
    if spool is None:
        raise ValueError('Bobina non trovata tra quelle eliminate.')
    wb['Bobine'].cell(spool['row'], 11).value = None
    sync_states(wb)


def material_group(material):
    material = material.strip().casefold()
    if material.startswith('abs'):
        return 'ABS'
    if material.startswith(('supporto', 'support')) or material in ('pva', 'bvoh'):
        return 'Supporto'
    return 'Altro'


def update_spool(wb, bid, material, brand, color, weight, used):
    b = next((b for b in bobine(wb, include_removed=True) if b['id'] == bid), None)
    if not b: raise ValueError('Bobina non trovata.')
    if not material.strip() or not color.strip(): raise ValueError('Inserisci materiale e colore.')
    if not all(math.isfinite(x) for x in (weight, used)) or weight <= 0 or not 0 <= used <= weight:
        raise ValueError('Il peso deve essere positivo e il consumo compreso tra zero e il peso iniziale.')
    recorded = sum(x['grammi'] for p in history(wb) for x in p['consumi'] if str(x['bobina']) == bid)
    if used < recorded - 1e-8:
        raise ValueError(f'Lo storico contiene già {recorded:g} g: correggi prima i consumi delle stampe.')
    ws = wb['Bobine']
    for col, value in enumerate([material.strip(), brand.strip(), color.strip(), weight, used], 2):
        ws.cell(b['row'], col, value)
    # The material correction applies to every historical use of this spool.
    for row in wb['Stampe'].iter_rows(min_row=2):
        if str(row[3].value) == bid: row[4].value = material.strip()
    sync_states(wb)


def update_print(wb, key, name, when, note, consumption, duration_minutes=None):
    original = next((p for p in history(wb) if p['key'] == key), None)
    if not original: raise ValueError('Stampa non trovata. Ricarica lo storico.')
    if not name.strip(): raise ValueError('Inserisci il nome della stampa.')
    if not isinstance(when, datetime): raise ValueError('Inserisci data e ora valide.')
    duration_minutes = original.get('durata') if duration_minutes is None else _duration(duration_minutes)
    bs = {b['id']: b for b in bobine(wb, include_removed=True)}
    old, new = defaultdict(float), defaultdict(float)
    for x in original['consumi']: old[str(x['bobina'])] += x['grammi']
    rows = []
    for x in consumption:
        try:
            g = float(x['grammi'])
            n = float(x['ugello'])
        except (ValueError, TypeError, KeyError):
            raise ValueError('Controlla ugello e grammi di ogni riga.')
        if not math.isfinite(g) or g < 0: raise ValueError('I grammi devono essere positivi o zero.')
        if n not in (1, 2, 3): raise ValueError('Scegli un ugello tra 1, 2 e 3.')
        if g == 0: continue
        bid = str(x.get('bobina') or '')
        if bid not in bs: raise ValueError('Scegli una bobina valida per ogni consumo.')
        rows.append([int(n), bid, bs[bid]['materiale'], g])
        new[bid] += g
    if not rows: raise ValueError('La stampa deve contenere almeno un consumo maggiore di zero.')
    updated = {}
    for bid in old.keys() | new.keys():
        if bid not in bs: raise ValueError(f'Bobina {bid} non trovata: impossibile correggere i consumi.')
        used = bs[bid]['usati'] - old[bid] + new[bid]
        if used < -1e-8 or used > bs[bid]['peso'] + 1e-8:
            raise ValueError(f'{bid}: la correzione produce un consumo totale di {used:g} g su {bs[bid]["peso"]:g} g. Controlla i valori.')
        updated[bid] = min(bs[bid]['peso'], max(0, used))
    # All checks finish before changing the workbook; other prints remain untouched.
    for bid, used in updated.items(): wb['Bobine'].cell(bs[bid]['row'], 6, used)
    ws = wb['Stampe']
    pid = ws.cell(original['rows'][0], 8).value or str(uuid4())
    for column, heading in PRINT_METADATA_HEADERS.items():
        ws.cell(1, column, heading)
    for index, use in enumerate(rows):
        values = [
            when, name.strip(), *use, note.strip(), pid, duration_minutes,
            original.get('planned_start'),
        ]
        if index < len(original['rows']):
            for c, value in enumerate(values, 1): ws.cell(original['rows'][index], c).value = value
        else: ws.append(values)
    for r in original['rows'][len(rows):]:
        for cell in ws[r]: cell.value = None
    sync_states(wb)


def delete_print(wb, key):
    original = next((p for p in history(wb) if p['key'] == key), None)
    if not original:
        raise ValueError('Stampa non trovata. Ricarica lo storico.')
    bs = {b['id']: b for b in bobine(wb, include_removed=True)}
    refunds = defaultdict(float)
    for use in original['consumi']:
        refunds[str(use['bobina'])] += use['grammi']
    for bid, grams in refunds.items():
        if bid not in bs:
            raise ValueError(f'Bobina {bid} non trovata: impossibile ripristinare i consumi.')
        if not math.isfinite(grams) or grams < 0 or bs[bid]['usati'] < grams - 1e-8:
            raise ValueError(f'{bid}: i consumi della bobina non corrispondono allo storico. Correggili prima di eliminare la stampa.')
    for bid, grams in refunds.items():
        wb['Bobine'].cell(bs[bid]['row'], 6, max(0, bs[bid]['usati'] - grams))
    # Clear records without shifting the row-based identifiers of older prints.
    for r in original['rows']:
        for cell in wb['Stampe'][r]:
            cell.value = None
    sync_states(wb)
