import tempfile
import unittest
import json
from datetime import datetime, time, timedelta
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest
import filament_store as store

ROOT = Path(__file__).resolve().parents[1]

class DashboardTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / 'test.xlsx'
        fixture = Workbook()
        fixture.active.title = 'Bobine'
        fixture['Bobine'].append(['ID', 'Materiale', 'Marca', 'Colore', 'Peso', 'Usati', 'Rimasti', '%', 'Stato', 'Note'])
        for row in [
            ['B001', 'ABS', '3ntr', 'Nero', 5000, 4900],
            ['B002', 'Supporto', '3ntr', 'Bianco', 1000, 300],
            ['B003', 'ABS', '3ntr', 'Arancione', 1000, 750],
        ]:
            fixture['Bobine'].append(row)
        fixture.create_sheet('Ugelli').append(['Ugello', 'Bobina'])
        for n in (1, 2, 3): fixture['Ugelli'].append([n, f'B{n:03d}'])
        fixture.create_sheet('Stampe').append(['Data', 'Nome', 'Ugello', 'Bobina', 'Materiale', 'Grammi', 'Note'])
        fixture['Stampe'].append([datetime(2026, 1, 1, 12), 'Stampa precedente', 1, 'B001', 'ABS', 100, 'Test storico'])
        fixture.save(self.path)
        self.wb = store.load(self.path)

    def tearDown(self): self.tmp.cleanup()

    def test_multinozzle_print_and_atomic_validation(self):
        before = {b['id']: b['rim'] for b in store.bobine(self.wb)}
        count = len(store.history(self.wb))
        with self.assertRaises(ValueError):
            store.record_print(self.wb, 'Invalid', {1: 1, 2: 999999, 3: 0})
        self.assertEqual(before, {b['id']: b['rim'] for b in store.bobine(self.wb)})
        store.record_print(self.wb, 'Test', {1: 1.5, 2: 2, 3: 3})
        self.assertEqual(len(store.history(self.wb)), count + 1)
        store.save(self.wb, self.path)
        latest = store.history(store.load(self.path))[0]
        self.assertEqual(latest['totale'], 6.5)
        self.assertEqual(len(latest['consumi']), 3)
        self.assertTrue(self.path.with_suffix('.backup.xlsx').exists())

    def test_planning_queue_calendar_and_completion(self):
        before = {b['id']: b['usati'] for b in store.bobine(self.wb)}
        first = store.add_planned_print(
            self.wb, 'Lavoro lungo', {1: 12.5, 2: 6, 3: 0}, 150,
            'Preparare il piano', created=datetime(2026, 1, 2, 9),
        )
        second = store.add_planned_print(
            self.wb, 'Lavoro breve', {1: 5, 2: 0, 3: 0}, 45,
            created=datetime(2026, 1, 2, 10),
        )
        self.assertEqual(before, {b['id']: b['usati'] for b in store.bobine(self.wb)})
        self.assertEqual(len(store.planned_prints(self.wb)), 2)
        first_plan = next(p for p in store.planned_prints(self.wb) if p['key'] == first)
        self.assertIsNone(first_plan['inizio'])
        self.assertEqual(first_plan['priorita'], 'Quando possibile')
        store.update_planned_priority(self.wb, first, 'urgente')
        self.assertEqual(next(p for p in store.planned_prints(self.wb) if p['key'] == first)['priorita'], 'Urgente')
        with self.assertRaisesRegex(ValueError, 'priorità'):
            store.update_planned_priority(self.wb, first, 'Non valida')

        start = datetime(2026, 1, 5, 8)
        store.schedule_planned_print(self.wb, first, start)
        with self.assertRaisesRegex(ValueError, 'sovrappone'):
            store.schedule_planned_print(self.wb, second, start + timedelta(hours=1))
        store.schedule_planned_print(self.wb, second, start + timedelta(minutes=150))
        store.unschedule_planned_print(self.wb, second)
        self.assertIsNone(next(p for p in store.planned_prints(self.wb) if p['key'] == second)['inizio'])

        store.complete_planned_print(self.wb, first, datetime(2026, 1, 5, 10, 30), [
            {'ugello': 1, 'bobina': 'B001', 'grammi': 10},
            {'ugello': 2, 'bobina': 'B002', 'grammi': 4},
        ])
        remaining = store.planned_prints(self.wb)
        self.assertEqual([p['key'] for p in remaining], [second])
        completed = store.history(self.wb)[0]
        self.assertEqual((completed['nome'], completed['durata'], completed['totale']), ('Lavoro lungo', 150, 14))
        self.assertEqual(completed['data'], datetime(2026, 1, 5, 10, 30))
        self.assertEqual(completed['planned_start'], start)
        self.assertEqual(completed['calendar_start'], start)
        after = {b['id']: b['usati'] for b in store.bobine(self.wb)}
        self.assertEqual(after['B001'], before['B001'] + 10)
        self.assertEqual(after['B002'], before['B002'] + 4)

        store.delete_planned_print(self.wb, second)
        self.assertEqual(store.planned_prints(self.wb), [])
        store.save(self.wb, self.path)
        loaded = store.load(self.path)
        self.assertIn('Pianificazione', loaded.sheetnames)
        self.assertEqual(store.history(loaded)[0]['durata'], 150)
        self.assertEqual(loaded['Stampe'].cell(1, 10).value, 'Inizio nel calendario')
        self.assertEqual(store.history(loaded)[0]['calendar_start'], start)

    def test_history_calendar_start_falls_back_to_actual_date_and_survives_edits(self):
        actual = datetime(2026, 1, 3, 14, 15)
        store.record_print(self.wb, 'Stampa diretta', {1: 5, 2: 0, 3: 0}, when=actual, duration_minutes=75)
        direct = store.history(self.wb)[0]
        self.assertIsNone(direct['planned_start'])
        self.assertEqual(direct['calendar_start'], actual)

        planned_start = datetime(2026, 1, 6, 8, 30)
        key = store.add_planned_print(
            self.wb, 'Stampa pianificata', {1: 5, 2: 0, 3: 0}, 90,
            start=planned_start,
        )
        store.complete_planned_print(self.wb, key, datetime(2026, 1, 6, 10, 15))
        planned = store.history(self.wb)[0]
        store.update_print(
            self.wb, planned['key'], 'Stampa pianificata corretta',
            datetime(2026, 1, 6, 10, 30), planned['note'], planned['consumi'], 100,
        )
        corrected = store.history(self.wb)[0]
        self.assertEqual(corrected['planned_start'], planned_start)
        self.assertEqual(corrected['calendar_start'], planned_start)
        self.assertEqual(corrected['durata'], 100)

    def test_planned_completion_is_atomic_when_stock_is_insufficient(self):
        key = store.add_planned_print(self.wb, 'Troppo materiale', {1: 50, 2: 0, 3: 0}, 60)
        before = {b['id']: b['usati'] for b in store.bobine(self.wb)}
        with self.assertRaisesRegex(ValueError, 'disponibili'):
            store.complete_planned_print(self.wb, key, consumption=[
                {'ugello': 1, 'bobina': 'B001', 'grammi': 101},
            ])
        self.assertEqual(before, {b['id']: b['usati'] for b in store.bobine(self.wb)})
        self.assertEqual(len(store.planned_prints(self.wb)), 1)

    def test_optimized_suggestion_respects_work_hours_and_turnaround(self):
        monday = datetime(2026, 1, 5, 8, 30)
        store.add_planned_print(self.wb, 'Attiva', {1: 5, 2: 0, 3: 0}, 90, start=monday)
        optimal = store.add_planned_print(self.wb, 'Ottimale', {1: 5, 2: 0, 3: 0}, 300)
        store.add_planned_print(self.wb, 'Troppo lunga', {1: 5, 2: 0, 3: 0}, 420)
        store.add_planned_print(self.wb, 'Breve', {1: 5, 2: 0, 3: 0}, 60)
        store.add_planned_print(self.wb, 'Già fissata', {1: 5, 2: 0, 3: 0}, 60,
                                start=datetime(2026, 1, 5, 16, 30))
        suggestion = store.suggest_next_print(store.planned_prints(self.wb), now=datetime(2026, 1, 5, 9))
        self.assertEqual(suggestion['proposta']['key'], optimal)
        self.assertEqual(suggestion['inizio'], datetime(2026, 1, 5, 10, 30))
        self.assertEqual(suggestion['fine'], datetime(2026, 1, 5, 15, 30))
        self.assertEqual(suggestion['prossima_programmata']['nome'], 'Già fissata')
        self.assertIsNone(store.suggest_next_print(store.planned_prints(self.wb), now=datetime(2026, 1, 5, 7)))

    def test_optimized_suggestion_prefers_priority_then_duration(self):
        monday = datetime(2026, 1, 5, 8, 30)
        store.add_planned_print(self.wb, 'Attiva', {1: 5, 2: 0, 3: 0}, 60, start=monday)
        store.add_planned_print(
            self.wb, 'Lunga ma normale', {1: 5, 2: 0, 3: 0}, 300,
            priority='Quando possibile',
        )
        urgent = store.add_planned_print(
            self.wb, 'Urgente', {1: 5, 2: 0, 3: 0}, 180, priority='Urgente',
        )
        immediate = store.add_planned_print(
            self.wb, 'Subito', {1: 5, 2: 0, 3: 0}, 60, priority='SUBITO',
        )
        plans = store.planned_prints(self.wb)
        suggestion = store.suggest_next_print(plans, now=datetime(2026, 1, 5, 9))
        self.assertEqual(suggestion['proposta']['key'], immediate)
        self.assertEqual([p['priorita'] for p in plans if p['inizio'] is None], [
            'SUBITO', 'Urgente', 'Quando possibile',
        ])

        store.update_planned_priority(self.wb, urgent, 'SUBITO')
        suggestion = store.suggest_next_print(store.planned_prints(self.wb), now=datetime(2026, 1, 5, 9))
        self.assertEqual(suggestion['proposta']['key'], urgent)

    def test_optimized_suggestion_continues_after_an_accepted_proposal(self):
        monday = datetime(2026, 1, 5, 8, 30)
        store.add_planned_print(self.wb, 'Attiva', {1: 5, 2: 0, 3: 0}, 60, start=monday)
        first = store.add_planned_print(
            self.wb, 'Prima proposta', {1: 5, 2: 0, 3: 0}, 240, priority='SUBITO',
        )
        second = store.add_planned_print(
            self.wb, 'Seconda proposta', {1: 5, 2: 0, 3: 0}, 120, priority='Urgente',
        )
        suggestion = store.suggest_next_print(store.planned_prints(self.wb), now=datetime(2026, 1, 5, 9))
        self.assertEqual(suggestion['proposta']['key'], first)
        store.schedule_planned_print(self.wb, first, suggestion['inizio'])

        follow_up = store.suggest_next_print(store.planned_prints(self.wb), now=datetime(2026, 1, 5, 9))
        self.assertEqual(follow_up['proposta']['key'], second)
        self.assertEqual(follow_up['precedente']['key'], first)
        self.assertEqual(follow_up['inizio'], datetime(2026, 1, 5, 14, 30))
        self.assertEqual(follow_up['fine'], datetime(2026, 1, 5, 16, 30))

    def test_setup_and_exhaustion(self):
        with self.assertRaises(ValueError): store.set_assignments(self.wb, {1:'B001',2:'B001',3:''})
        amount = store.bobine(self.wb)[0]['rim']
        store.record_print(self.wb, 'Finish', {1:amount,2:0,3:0})
        self.assertEqual(store.bobine(self.wb)[0]['rim'], 0)
        with self.assertRaises(ValueError): store.set_assignments(self.wb, {1:'B001',2:'',3:''})
        store.set_assignments(self.wb, {1:'',2:'',3:''})
        self.assertTrue(all(not b['ugello'] for b in store.bobine(self.wb)))

    def test_purchase_clears_restock_and_conflicts_rejected(self):
        self.assertTrue(any(x['Materiale'] == 'ABS' and x['Colore'] == 'Nero' for x in store.restock(store.bobine(self.wb))))
        store.add_spool(self.wb, 'ABS', '3ntr', 'Nero', 5000)
        self.assertFalse(any(x['Materiale'] == 'ABS' and x['Colore'] == 'Nero' for x in store.restock(store.bobine(self.wb))))
        stale = store.load(self.path)
        store.save(self.wb, self.path)
        with self.assertRaises(ValueError): store.save(stale, self.path)

    def test_delete_restore_preserves_history_and_ids(self):
        before = store.history(self.wb)
        with self.assertRaises(ValueError): store.remove_spool(self.wb, 'B001')
        store.set_assignments(self.wb, {1:'', 2:'B002', 3:'B003'})
        store.remove_spool(self.wb, 'B001')
        self.assertNotIn('B001', [b['id'] for b in store.bobine(self.wb)])
        self.assertEqual(store.history(self.wb), before)
        archived = next(b for b in store.bobine(self.wb, include_removed=True) if b['id'] == 'B001')
        self.assertEqual(archived['colore'], 'Nero')
        with self.assertRaises(ValueError): store.set_assignments(self.wb, {1:'B001', 2:'', 3:''})
        bid = store.add_spool(self.wb, 'PLA', 'Test', 'Blu', 1000)
        store.remove_spool(self.wb, bid)
        new_id = store.add_spool(self.wb, 'PLA', 'Test', 'Verde', 1000)
        self.assertNotEqual(bid, new_id)
        store.save(self.wb, self.path)
        wb = store.load(self.path)
        store.restore_spool(wb, 'B001')
        self.assertIn('B001', [b['id'] for b in store.bobine(wb)])
        self.assertEqual(store.history(wb), before)

    def test_inventory_controls(self):
        bid = store.add_spool(self.wb, 'PLA', 'Test', 'Blu', 1000)
        store.save(self.wb, self.path)
        original_load, original_save = store.load, store.save
        with patch.object(store, 'FILE', self.path), patch.object(store, 'load', side_effect=lambda: original_load(self.path)), patch.object(store, 'save', side_effect=lambda wb: original_save(wb, self.path)):
            app = AppTest.from_file(str(ROOT / 'filament_dashboard.py')).run()
            self.assertNotIn('Riacquisti', app.sidebar.radio[0].options)
            app.sidebar.radio[0].set_value('Magazzino').run()
            for option in ['Grammi decrescenti', 'Colore A–Z', 'ID bobina']:
                next(x for x in app.selectbox if x.label == 'Ordina per').set_value(option).run()
                self.assertFalse(app.exception)
            self.assertTrue(app.button(key='delete_B001').disabled)
            app.button(key=f'delete_{bid}').click().run()
            self.assertFalse(app.exception)
            self.assertNotIn(bid, [b['id'] for b in store.bobine(original_load(self.path))])
            app.button(key=f'restore_{bid}').click().run()
            self.assertFalse(app.exception)
            self.assertIn(bid, [b['id'] for b in store.bobine(original_load(self.path))])

    def test_new_spool_guided_choices_and_custom_values(self):
        original_load, original_save = store.load, store.save
        with patch.object(store, 'FILE', self.path), \
             patch.object(store, 'load', side_effect=lambda: original_load(self.path)), \
             patch.object(store, 'save', side_effect=lambda wb: original_save(wb, self.path)):
            app = AppTest.from_file(str(ROOT / 'filament_dashboard.py')).run()
            app.sidebar.radio[0].set_value('Magazzino').run()
            self.assertEqual(app.selectbox(key='new_material_choice').value, 'ABS')
            self.assertEqual(app.selectbox(key='new_brand_choice').value, '3ntr')
            self.assertEqual(app.selectbox(key='new_color_choice').value, 'Nero')
            app.selectbox(key='new_material_choice').set_value('Altro').run()
            app.selectbox(key='new_brand_choice').set_value('Altro').run()
            app.selectbox(key='new_color_choice').set_value('Altro').run()
            app.text_input(key='new_material_custom').set_value('PETG')
            app.text_input(key='new_brand_custom').set_value('Prusament')
            app.text_input(key='new_color_custom').set_value('Galaxy Black')
            app.number_input(key='new_spool_weight').set_value(750)
            app.button(key='add_new_spool').click().run()
            self.assertFalse(app.exception)
            added = store.bobine(original_load(self.path))[-1]
            self.assertEqual(
                (added['materiale'], added['marca'], added['colore'], added['peso']),
                ('PETG', 'Prusament', 'Galaxy Black', 750),
            )
            self.assertEqual(app.selectbox(key='new_material_choice').value, 'ABS')
            self.assertEqual(app.selectbox(key='new_brand_choice').value, '3ntr')

    def test_edit_print_moves_consumption_and_rejects_overdraw(self):
        p = store.history(self.wb)[0]
        before = {b['id']: b['usati'] for b in store.bobine(self.wb)}
        with self.assertRaises(ValueError):
            store.update_print(self.wb, p['key'], 'Invalid', p['data'], '', [{'ugello': 2, 'bobina': 'B002', 'grammi': 2000}])
        self.assertEqual(before, {b['id']: b['usati'] for b in store.bobine(self.wb)})
        store.update_print(self.wb, p['key'], 'Corretta', p['data'], 'Note nuove', [
            {'ugello': 2, 'bobina': 'B002', 'grammi': 20},
            {'ugello': 3, 'bobina': 'B003', 'grammi': 5}])
        bs = {b['id']: b for b in store.bobine(self.wb)}
        self.assertEqual(bs['B001']['usati'], before['B001'] - 100)
        self.assertEqual(bs['B002']['usati'], before['B002'] + 20)
        updated = store.history(self.wb)[0]
        self.assertEqual(updated['totale'], 25)
        self.assertEqual(updated['nome'], 'Corretta')
        store.record_print(self.wb, 'Altra stampa', {1: 1, 2: 0, 3: 0})
        store.update_print(self.wb, updated['key'], 'Corretta', p['data'], '', [
            {'ugello': 2, 'bobina': 'B002', 'grammi': 10}])
        store.save(self.wb, self.path)
        result = store.history(store.load(self.path))
        self.assertEqual(len(result), 2)
        self.assertEqual(next(x for x in result if x['nome'] == 'Corretta')['totale'], 10)
        self.assertEqual(next(x for x in result if x['nome'] == 'Altra stampa')['totale'], 1)

    def test_edit_spool_validation_and_history_material(self):
        with self.assertRaises(ValueError): store.update_spool(self.wb, 'B001', 'PLA', 'Test', 'Blu', 50, 100)
        with self.assertRaises(ValueError): store.update_spool(self.wb, 'B001', 'PLA', 'Test', 'Blu', 5000, 50)
        store.update_spool(self.wb, 'B001', 'PLA', 'Test', 'Blu', 6000, 4900)
        b = next(b for b in store.bobine(self.wb) if b['id'] == 'B001')
        self.assertEqual(b['rim'], 1100)
        self.assertEqual(b['colore'], 'Blu')
        self.assertEqual(store.history(self.wb)[0]['consumi'][0]['materiale'], 'PLA')
        self.assertEqual(store.assignments(self.wb)[1], 'B001')

    def test_edit_dialogs_save_and_cancel(self):
        original_load, original_save = store.load, store.save
        with patch.object(store, 'FILE', self.path), patch.object(store, 'load', side_effect=lambda: original_load(self.path)), patch.object(store, 'save', side_effect=lambda wb: original_save(wb, self.path)):
            app = AppTest.from_file(str(ROOT / 'filament_dashboard.py')).run()
            app.sidebar.radio[0].set_value('Magazzino').run()
            app.button(key='edit_spool_B001').click().run()
            self.assertFalse(app.exception)
            app.text_input(key='edit_color').set_value('Grigio')
            next(x for x in app.button if x.label == 'Salva modifiche').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(store.bobine(original_load(self.path))[0]['colore'], 'Grigio')
            app.sidebar.radio[0].set_value('Storico').run()
            next(x for x in app.button if x.key and x.key.startswith('edit_print_')).click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.dataframe[0].value.iloc[0]['bobina'], 'B001 · ABS / Grigio')
            app.text_input(key='edit_print_name').set_value('Nome corretto')
            next(x for x in app.button if x.label == 'Salva modifiche').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(store.history(original_load(self.path))[0]['nome'], 'Nome corretto')
            next(x for x in app.button if x.key and x.key.startswith('edit_print_')).click().run()
            app.text_input(key='edit_print_name').set_value('Da annullare')
            next(x for x in app.button if x.label == 'Annulla').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(store.history(original_load(self.path))[0]['nome'], 'Nome corretto')

    def test_delete_print_refunds_all_nozzles(self):
        before = {b['id']: b['usati'] for b in store.bobine(self.wb)}
        original = store.history(self.wb)
        store.record_print(self.wb, 'Da eliminare', {1: 10, 2: 20, 3: 30})
        p = store.history(self.wb)[0]
        store.delete_print(self.wb, p['key'])
        self.assertEqual(before, {b['id']: b['usati'] for b in store.bobine(self.wb)})
        self.assertEqual(original, store.history(self.wb))
        with self.assertRaises(ValueError): store.delete_print(self.wb, p['key'])
        self.assertEqual(before, {b['id']: b['usati'] for b in store.bobine(self.wb)})

    def test_delete_print_button(self):
        original_load, original_save = store.load, store.save
        with patch.object(store, 'FILE', self.path), patch.object(store, 'load', side_effect=lambda: original_load(self.path)), patch.object(store, 'save', side_effect=lambda wb: original_save(wb, self.path)):
            app = AppTest.from_file(str(ROOT / 'filament_dashboard.py')).run()
            app.sidebar.radio[0].set_value('Storico').run()
            app.button(key='delete_print_legacy-2').click().run()
            self.assertFalse(app.exception)
            self.assertTrue(app.success)
            self.assertEqual(store.history(original_load(self.path)), [])
            self.assertEqual(store.bobine(original_load(self.path))[0]['usati'], 4800)

    def test_overview_shows_active_and_next_print(self):
        now = datetime.now().replace(second=0, microsecond=0)
        store.add_planned_print(
            self.wb, 'Stampa attiva test', {1: 5, 2: 0, 3: 0}, 60,
            start=now - timedelta(minutes=15), priority='SUBITO',
        )
        store.add_planned_print(
            self.wb, 'Stampa successiva test', {1: 5, 2: 0, 3: 0}, 90,
            start=now + timedelta(minutes=60), priority='Urgente',
        )
        store.save(self.wb, self.path)
        original_load, original_save = store.load, store.save
        with patch.object(store, 'FILE', self.path), \
             patch.object(store, 'load', side_effect=lambda: original_load(self.path)), \
             patch.object(store, 'save', side_effect=lambda wb: original_save(wb, self.path)):
            app = AppTest.from_file(str(ROOT / 'filament_dashboard.py')).run()
            self.assertFalse(app.exception)
            self.assertEqual(app.subheader[0].value, 'Produzione')
            markup = ' '.join(block.value for block in app.markdown)
            self.assertIn('STAMPA IN CORSO', markup)
            self.assertIn('Stampa attiva test', markup)
            self.assertIn('STAMPA SUCCESSIVA', markup)
            self.assertIn('Stampa successiva test', markup)

    def test_all_pages_and_print_submission(self):
        original_load, original_save = store.load, store.save
        with patch.object(store, 'FILE', self.path), patch.object(store, 'load', side_effect=lambda: original_load(self.path)), patch.object(store, 'save', side_effect=lambda wb: original_save(wb, self.path)):
            app = AppTest.from_file(str(ROOT / 'filament_dashboard.py')).run()
            self.assertFalse(app.exception)
            for page in ['Magazzino','Storico','Pianificazione','Nuova stampa']:
                app.sidebar.radio[0].set_value(page).run()
                self.assertFalse(app.exception, page)
            app.text_input(key='print_name').set_value('UI test')
            app.number_input(key='cons_1_B001').set_value(2.5)
            next(b for b in app.button if b.label == 'Registra stampa e aggiorna scorte').click().run()
            self.assertFalse(app.exception)
            self.assertTrue(app.success)
            self.assertEqual(app.sidebar.radio[0].value, 'Panoramica')
            self.assertIn('Stampa salvata', app.success[0].value)
            app.sidebar.radio[0].set_value('Nuova stampa').run()
            self.assertEqual(app.number_input(key='cons_1_B001').value, 0)
            self.assertEqual(app.text_input(key='print_name').value, '')
            self.assertTrue(any(p['nome'] == 'UI test' for p in store.history(original_load(self.path))))

    def test_new_print_can_be_sent_to_planning_queue(self):
        original_load, original_save = store.load, store.save
        with patch.object(store, 'FILE', self.path), patch.object(store, 'load', side_effect=lambda: original_load(self.path)), patch.object(store, 'save', side_effect=lambda wb: original_save(wb, self.path)):
            app = AppTest.from_file(str(ROOT / 'filament_dashboard.py')).run()
            app.sidebar.radio[0].set_value('Nuova stampa').run()
            next(x for x in app.radio if x.label == 'Stato della stampa').set_value('Da programmare').run()
            app.text_input(key='print_name').set_value('Da calendarizzare')
            app.radio(key='print_priority').set_value('SUBITO')
            app.number_input(key='cons_1_B001').set_value(8)
            next(b for b in app.button if b.label == 'Aggiungi alla coda di pianificazione').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.sidebar.radio[0].value, 'Pianificazione')
            self.assertEqual(
                [heading.value for heading in app.subheader],
                ['Stampe da pianificare', 'Calendario settimanale', 'Stampe in calendario'],
            )
            queued = store.planned_prints(original_load(self.path))
            self.assertEqual(len(queued), 1)
            self.assertEqual(
                (queued[0]['nome'], queued[0]['durata'], queued[0]['inizio'], queued[0]['priorita']),
                ('Da calendarizzare', 60, None, 'SUBITO'),
            )
            self.assertEqual(store.bobine(original_load(self.path))[0]['usati'], 4900)

            app.button(key=f'priority_plan_{queued[0]["key"]}').click().run()
            app.radio(key=f'priority_value_{queued[0]["key"]}').set_value('Urgente')
            next(b for b in app.button if b.label == 'Salva priorità').click().run()
            self.assertFalse(app.exception)
            self.assertEqual(store.planned_prints(original_load(self.path))[0]['priorita'], 'Urgente')

            app.button(key=f'schedule_{queued[0]["key"]}').click().run()
            self.assertFalse(app.exception)
            app.date_input(key=f'schedule_day_{queued[0]["key"]}').set_value(datetime.now().date() + timedelta(days=1))
            app.time_input(key=f'schedule_time_{queued[0]["key"]}').set_value(time(9, 0))
            next(b for b in app.button if b.label == 'Salva nel calendario').click().run()
            self.assertFalse(app.exception)
            scheduled = store.planned_prints(original_load(self.path))[0]
            self.assertEqual(scheduled['inizio'].time(), time(9, 0))

            calendar_key = app.get('component_instance')[0].key
            app.button(key=f'unschedule_plan_{scheduled["key"]}').click().run()
            self.assertFalse(app.exception)
            self.assertIsNone(store.planned_prints(original_load(self.path))[0]['inizio'])
            self.assertEqual(len(app.get('component_instance')), 1)
            self.assertNotEqual(app.get('component_instance')[0].key, calendar_key)

            app.button(key=f'schedule_{scheduled["key"]}').click().run()
            app.date_input(key=f'schedule_day_{scheduled["key"]}').set_value(datetime.now().date() + timedelta(days=1))
            app.time_input(key=f'schedule_time_{scheduled["key"]}').set_value(time(9, 0))
            next(b for b in app.button if b.label == 'Salva nel calendario').click().run()
            self.assertFalse(app.exception)
            scheduled = store.planned_prints(original_load(self.path))[0]

            app.button(key=f'complete_plan_{scheduled["key"]}').click().run()
            self.assertFalse(app.exception)
            next(b for b in app.button if b.label == 'Completa e aggiorna scorte').click().run()
            self.assertFalse(app.exception)
            completed_book = original_load(self.path)
            self.assertEqual(store.planned_prints(completed_book), [])
            self.assertEqual(store.history(completed_book)[0]['nome'], 'Da calendarizzare')
            self.assertEqual(store.bobine(completed_book)[0]['usati'], 4908)
            calendar_args = json.loads(app.get('component_instance')[0].proto.json_args)
            completed_events = [
                event for event in calendar_args['events']
                if event.get('extendedProps', {}).get('planningState') == 'completed'
            ]
            self.assertEqual(len(completed_events), 1)
            completed_event = next(event for event in completed_events if 'Da calendarizzare' in event['title'])
            self.assertFalse(completed_event['editable'])
            self.assertIn('completed-event', completed_event['classNames'])

if __name__ == '__main__': unittest.main()
